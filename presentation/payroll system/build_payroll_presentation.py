from pathlib import Path
import sys

PRESENTATION_DIR = Path(__file__).resolve().parent
ROOT = PRESENTATION_DIR.parent.parent
sys.path.insert(0, str(ROOT / '.tmp-pptx-libs'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NAVY, TEAL, WHITE, MUTED, PANEL = '102338', '42D6B1', 'FFFFFF', 'B8C7D8', '1C354E'

def box(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape

def text(slide, x, y, w, h, value, size=22, color=WHITE, bold=False):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = 'Aptos'; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(12)
    return tf

def slide(title, subtitle=''):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string(NAVY)
    text(s, .65, .28, 11, .35, 'PAYROLL  /  ATTENDANCE  /  DRIVER EARNINGS', 11, TEAL, True)
    text(s, .65, .95, 12, .85, title, 34, WHITE, True)
    if subtitle: text(s, .65, 1.88, 12, .7, subtitle, 18, MUTED)
    text(s, .65, 7.03, 11.4, .25, 'PROPOSED SCOPE & COMMERCIAL TERMS  •  FOR DISCUSSION', 10, MUTED)
    text(s, 12.1, 6.99, .6, .35, f'{len(prs.slides):02}', 12, TEAL)
    return s

def card(s, x, title, body, width=5.85):
    box(s, x, 2.85, width, 3.7, PANEL)
    text(s, x+.25, 3.08, width-.5, .6, title, 25, TEAL, True)
    text(s, x+.25, 3.85, width-.5, 2.5, body, 21)

def table(s, headers, rows, widths, top=2.85, row_height=.55, font=19):
    x = .65
    for i, row in enumerate([headers]+rows):
        box(s, .65, top+i*row_height, 12.0, row_height-.025, PANEL if i%2 == 0 else NAVY)
        x=.65
        for value, width in zip(row, widths):
            text(s, x+.15, top+i*row_height+.07, width-.25, row_height-.1, str(value), font, TEAL if i == 0 else WHITE, i == 0)
            x += width

s=slide('Payroll, attendance & driver earnings', 'A proposed system tailored to your company’s payroll workflow.')
card(s,.65,'SYSTEM SCOPE','Payroll and employee records\nAttendance and leave\nDriver-task earnings\nHR review, payslips and reports')
card(s,6.8,'TODAY’S PRESENTATION','What the system will cover\nImplementation and ongoing services\nPayment options\nMonthly costs over 1, 2, 3 and 5 years')

s=slide('Bring payroll preparation together', 'Help HR move from scattered inputs to reviewed, approved pay.')
card(s,.65,'ONE CONNECTED WORKFLOW','Employee records and pay rates\nAttendance, leave and deductions\nAccepted driver-task earnings')
card(s,6.8,'CLEARER REVIEW','Trace how each pay amount was calculated\nReview before finalization\nIssue payslips and retain payroll history')

s=slide('Four parts of the proposed solution')
table(s,['Component','Proposed coverage'],[
['Payroll & HR','Employees, semi-monthly payroll, deductions and reports'],
['Time-in / Time-out','Browser attendance, hours, lateness and HR corrections'],
['Driver-task integration','Eligible completions, task rates and payment checks'],
['Deployment & training','Setup, initial data assistance and HR documentation']], [3.6,8.4],row_height=.75,font=19)
text(s,.8,6.15,11.8,.55,'Company policies and statutory calculations are validated with HR/accounting before acceptance.',16,MUTED)

s=slide('Payroll & HR — proposed scope')
card(s,.65,'EMPLOYEE & PAYROLL RECORDS','Employee profiles and pay rates\nSemi-monthly payroll processing\nAllowances, incentives and deductions\nLoans and cash advances')
card(s,6.8,'CALCULATIONS & REVIEW','Government contributions and tax\nOvertime, premiums and night pay\n13th-month and final-pay preparation\nApprovals, payslips and reports')
text(s,.8,6.58,11.8,.3,'Calculation policies and acceptance examples are agreed with HR/accounting.',14,MUTED)

s=slide('Attendance & driver earnings — proposed scope')
card(s,.65,'ATTENDANCE & LEAVE','Browser-based time-in / time-out\nHours, lateness and undertime\nSchedules, holidays and leave\nHR corrections before payroll')
card(s,6.8,'DRIVER EARNINGS','Match drivers to payroll employees\nUse eligible accepted completions\nApply agreed task rates\nPrevent duplicate task payments')

s=slide('What else is covered')
card(s,.65,'IMPLEMENTATION','Requirements and policy confirmation\nConfiguration and initial data assistance\nPayroll and integration validation\nDeployment, HR training and guides')
card(s,6.8,'RECORDS & CONTROLS','Authorized user roles\nAudit history for payroll changes\nProtected finalized payroll records\nTask reconciliation and data exports')

s=slide('Service coverage by payment option')
card(s,.65,'ONE-TIME PURCHASE','90-day defect warranty\nMinor updates within existing retainer terms\nHosting funded separately\nMajor enhancements quoted separately')
card(s,6.8,'MONTHLY SUBSCRIPTION','Payroll hosting and backups\nStandard support and defect corrections\nRoutine supported statutory updates\nRequested custom changes billed separately')
text(s,.8,6.58,11.8,.3,'Biometrics, external service fees and new integrations are separately quoted.',14,MUTED)

s=slide('The difference: how changes are covered', 'Both options deliver the requirements agreed at the start.')
card(s,.65,'ONE-TIME + EXISTING RETAINER','Agreed project changes are implemented within the approved scope.\nMinor adjustments and updates are included within the ₱10,000 retainer coverage.')
card(s,6.8,'MONTHLY SUBSCRIPTION','Requested customizations, workflow changes and additional features receive a separate quotation.\nWork begins after the quotation is approved.')
s.notes_slide.notes_text_frame.text='The retainer billing frequency and exact payroll coverage have not been confirmed. Do not describe it as a monthly fee. Defect corrections are distinct from requested custom changes.'

s=slide('Option 1 — Invest once for permanent use','₱320,000 project price • No recurring per-employee software fee')
card(s,.65,'LONG-TERM VALUE','Permanent use by one company\nSoftware fee stays fixed as the workforce grows\nPayments follow delivery milestones')
card(s,6.8,'CHANGES & SUPPORT','90-day defect warranty after acceptance\nMinor updates under the existing ₱10,000 retainer\nMajor enhancements quoted separately')

s=slide('What the ₱320,000 project covers')
table(s,['Component','Amount'],[['Payroll & HR system','₱250,000'],['Time-in / Time-out','₱20,000'],['Driver-task integration','₱50,000'],['Deployment, training & documentation','Included'],['TOTAL PROJECT PRICE','₱320,000']],[9,3],row_height=.62)
text(s,.8,6.25,11.8,.5,'Deployment, training and documentation are included. Existing retainer terms are separate.',16,MUTED)

s=slide('Option 2 — Start with a smaller investment','₱35,000 setup + ₱85 per active employee/month • ₱3,000 monthly minimum')
table(s,['Active employees','Monthly fee','First year incl. setup'],[['25','₱3,000','₱71,000'],['50','₱4,250','₱86,000'],['100','₱8,500','₱137,000'],['200','₱17,000','₱239,000'],['300','₱25,500','₱341,000']],[4,4,4],row_height=.55)
text(s,.8,6.25,11.8,.5,'Assumes 12 months and a constant headcount. Each active person is counted once per month.',15,MUTED)

s=slide('Monthly plan — costs over time','Totals include the ₱35,000 setup fee once, plus monthly subscription payments.')
years = (1, 2, 3, 5)
cost_rows = []
for employees in (25, 50, 100, 200, 300):
    monthly = max(3000, employees * 85)
    cost_rows.append([str(employees)] + [f'₱{35000 + monthly * 12 * year:,}' for year in years])
table(s,['Employees','1 year','2 years','3 years','5 years'],cost_rows,[2.2,2.45,2.45,2.45,2.45],row_height=.55,font=18)
text(s,.8,6.23,11.8,.55,'Assumes constant headcount and rate. Excludes custom changes, external fees and applicable taxes.',15,MUTED)
s.notes_slide.notes_text_frame.text='Formula: PHP 35,000 + max(PHP 3,000, employee count × PHP 85) × 12 × years. These are cumulative subscription totals, not annual prepaid packages. One-time ownership also has retainer and operating costs; do not compare these figures as total ownership costs.'

s=slide('Delivery milestones for the one-time option')
table(s,['Milestone','Share','Payment'],[['Commencement & requirements approval','30%','₱96,000'],['Core HR, attendance & payroll demonstration','30%','₱96,000'],['Integration & user acceptance testing','30%','₱96,000'],['Deployment & final acceptance','10%','₱32,000']],[8,1.5,2.5],row_height=.72,font=18)
text(s,.8,6.15,11.8,.6,'Monthly option: setup payable at commencement; subscription starts at production activation.',17,MUTED)

s=slide('Agree the details. Validate. Launch.','Confirm the scope and payment model with HR and management.')
card(s,.65,'DECISIONS TO CONFIRM','Payroll policies and employee count\nMinor-update limits and retainer coverage\nHosting and backup responsibilities\nSchedule and acceptance criteria')
card(s,6.8,'COMMERCIAL BOUNDARIES','Major changes quoted separately\nOne-time hosting costs funded separately\nDeveloper retains reusable source code\nTaxes and final terms confirmed in the quotation')
s.notes_slide.notes_text_frame.text='Retainer frequency is deliberately unspecified. The one-time purchase grants company use, not source-code ownership. Do not promise unlimited changes or free hosting. Proposed scope is subject to agreement.'

s=slide('System inclusions', 'What is included in the agreed payroll system.')
card(s,.65,'PAYROLL & ATTENDANCE','Employee records & payroll — Maintain employee details and calculate semi-monthly pay.\nContributions & tax — Calculate SSS, PhilHealth, Pag-IBIG and withholding tax.\nAttendance & leave — Record time-in/out, worked hours and leave for HR review.')
card(s,6.8,'EARNINGS, REPORTS & SETUP','Driver-task earnings — Credit eligible completed tasks at agreed rates and prevent duplicate payments.\nPayslips & reports — Issue pay breakdowns, payroll summaries and task reconciliation.\nSetup & training — Configure and deploy the system, assist initial data setup and provide training and guides.')
s.notes_slide.notes_text_frame.text='Includes initial employee-data assistance and configured payroll reports/CSV or print exports. Government contribution and tax calculations do not imply government-portal submission or every filing format. Agreed test cutoffs must reconcile with HR-verified calculations. Training sessions, initial import volume, task eligibility/cutoff rules and retainer coverage must be confirmed in the final scope. Browser attendance does not promise offline attendance. One-time hosting is client-funded; subscription payroll hosting and backups are included. The existing PHP 10,000 retainer frequency remains unconfirmed.'

output=PRESENTATION_DIR/'payroll-proposal-v2.pptx'
prs.save(output)
reopened=Presentation(output)
assert len(reopened.slides)==15
for s in reopened.slides:
    for shape in s.shapes:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left+shape.width <= reopened.slide_width
        assert shape.top+shape.height <= reopened.slide_height
print(f'Created and validated {len(reopened.slides)} editable slides: {output}')
